from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "submission_attachments"
ASSET_DIR = OUT_DIR / "assets"
DOCX_PATH = OUT_DIR / "P0684 - Kepang AI Decision Intelligence untuk Resiliensi Ketahanan Pangan Daerah.docx"
PDF_NAME = "P0684 - Kepang AI Decision Intelligence untuk Resiliensi Ketahanan Pangan Daerah.pdf"

BLUE = "2E74B5"
DARK = "1F2937"
MUTED = "6B7280"
GREEN = "047857"
AMBER = "92400E"
RED = "991B1B"
LIGHT_BLUE = "E8EEF5"
LIGHT_GREEN = "ECFDF5"
LIGHT_AMBER = "FFFBEB"
LIGHT_GRAY = "F3F4F6"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=DARK, size=8.5):
    cell.text = ""
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = para.add_run(str(text))
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_width(cell, width_in):
    width = int(width_in * 1440)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row in table.rows:
        for idx, width in enumerate(widths):
            set_cell_width(row.cells[idx], width)

    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(int(w * 1440) for w in widths)))
    tbl_w.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width * 1440)))
        grid.append(col)


def add_table(doc, headers, rows, widths, font_size=8.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.allow_autofit = False
    hdr = table.rows[0].cells
    for i, label in enumerate(headers):
        set_cell_text(hdr[i], label, bold=True, color="111827", size=8.4)
        set_cell_shading(hdr[i], LIGHT_BLUE)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=font_size)
    set_table_widths(table, widths)
    doc.add_paragraph()
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    p.add_run(text)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.style = "Normal"
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
    p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_callout(doc, title, body, fill=LIGHT_GREEN, color=GREEN):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    para = cell.paragraphs[0]
    run = para.add_run(title)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.size = Pt(10)
    para.add_run("\n" + body)
    for run in para.runs:
        run.font.name = "Calibri"
    set_table_widths(table, [6.35])
    doc.add_paragraph()


def configure_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 14, 7),
        ("Heading 2", 13, BLUE, 10, 5),
        ("Heading 3", 11.5, "1F4D78", 7, 3),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    for name in ["List Bullet", "List Number"]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(10)
        style.paragraph_format.space_after = Pt(3)


def get_font(size=28, bold=False):
    candidates = [
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def draw_rounded(draw, xy, fill, outline="#CBD5E1", radius=18, width=2):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def wrap_text(draw, text, font, max_width):
    words = text.split()
    lines = []
    cur = ""
    for word in words:
        trial = f"{cur} {word}".strip()
        if draw.textbbox((0, 0), trial, font=font)[2] <= max_width:
            cur = trial
        else:
            if cur:
                lines.append(cur)
            cur = word
    if cur:
        lines.append(cur)
    return lines


def draw_box(draw, xy, title, subtitle, fill, accent):
    x1, y1, x2, y2 = xy
    draw_rounded(draw, xy, fill=fill, outline="#CBD5E1")
    draw.rectangle((x1, y1, x1 + 9, y2), fill=accent)
    title_font = get_font(24, bold=True)
    sub_font = get_font(18)
    draw.text((x1 + 24, y1 + 18), title, fill="#111827", font=title_font)
    y = y1 + 52
    for line in wrap_text(draw, subtitle, sub_font, x2 - x1 - 45)[:3]:
        draw.text((x1 + 24, y), line, fill="#475569", font=sub_font)
        y += 24


def arrow(draw, start, end, color="#64748B"):
    draw.line((start, end), fill=color, width=4)
    ex, ey = end
    sx, sy = start
    if abs(ex - sx) >= abs(ey - sy):
        direction = 1 if ex > sx else -1
        points = [(ex, ey), (ex - 16 * direction, ey - 9), (ex - 16 * direction, ey + 9)]
    else:
        direction = 1 if ey > sy else -1
        points = [(ex, ey), (ex - 9, ey - 16 * direction), (ex + 9, ey - 16 * direction)]
    draw.polygon(points, fill=color)


def create_architecture_image(path):
    img = Image.new("RGB", (1800, 1180), "white")
    draw = ImageDraw.Draw(img)
    title_font = get_font(34, bold=True)
    label_font = get_font(17, bold=True)
    note_font = get_font(17)

    draw.text((60, 35), "System Architecture Kepang AI", fill="#0F172A", font=title_font)
    draw.text(
        (60, 80),
        "Arsitektur berlapis dari sumber data hingga rekomendasi keputusan, dengan status data yang jelas dan dapat diaudit.",
        fill="#475569",
        font=get_font(20),
    )

    def section_label(text, y):
        draw.text((60, y), text.upper(), fill="#64748B", font=label_font)

    def spine(y1, y2):
        arrow(draw, (900, y1), (900, y2), color="#475569")

    section_label("Sumber Data", 130)
    source_boxes = [
        ((65, 160, 365, 285), "Harga Pangan", "BI Harga Pangan / Bapanas untuk pemantauan komoditas", "#EFF6FF", "#2563EB"),
        ((400, 160, 700, 285), "Cuaca BMKG", "Prakiraan desa dan sinyal risiko gangguan panen", "#ECFDF5", "#0F766E"),
        ((735, 160, 1035, 285), "Makro BPS/BI", "Inflasi, kurs, impor, dan indikator tekanan eksternal", "#FFF7ED", "#EA580C"),
        ((1070, 160, 1370, 285), "Google Routes", "ETA dan jarak rute jalan bila API key tersedia", "#F5F3FF", "#7C3AED"),
        ((1405, 160, 1705, 285), "Data Mitra", "Stok, gudang, carrier, kapasitas, dan biaya aktual", "#FEF2F2", "#DC2626"),
    ]
    for box in source_boxes:
        draw_box(draw, *box)

    spine(300, 335)

    section_label("Ingestion", 335)
    draw_box(
        draw,
        (210, 365, 1590, 470),
        "Data Ingestion & Normalisasi",
        "Scraper/API scheduler, validasi format, mapping wilayah, timestamp sumber, quality check, dan penyimpanan metadata.",
        "#F8FAFC",
        "#334155",
    )

    spine(485, 520)

    section_label("Core Platform", 520)
    core_boxes = [
        ((65, 550, 430, 695), "Frontend Web", "React + Vite untuk Dashboard, Evidence Room, forecasting, dan rekomendasi logistik", "#EFF6FF", "#2563EB"),
        ((500, 550, 865, 695), "Backend API", "Node.js + Express untuk orkestrasi data, scoring risiko, dan rekomendasi", "#F0FDF4", "#16A34A"),
        ((935, 550, 1300, 695), "Data Store", "PostgreSQL/Supabase untuk data operasional, metadata sumber, dan seed MVP", "#F1F5F9", "#334155"),
        ((1370, 550, 1735, 695), "Model & Lineage", "Forecasting, scenario shock rupiah, label real-time, official-release, forecast, unavailable", "#F5F3FF", "#7C3AED"),
    ]
    for box in core_boxes:
        draw_box(draw, *box)

    spine(710, 745)

    section_label("Output Keputusan", 745)
    output_boxes = [
        ((65, 775, 365, 910), "Resilience Score", "Prioritas wilayah dan komoditas berdasarkan gabungan risiko", "#F0FDF4", "#16A34A"),
        ((400, 775, 700, 910), "Scenario Planning", "Simulasi kurs, cuaca, supply-demand, dan biaya logistik", "#FFF7ED", "#EA580C"),
        ((735, 775, 1035, 910), "Action Cards", "Saran intervensi stok, operasi pasar, dan mitigasi panen", "#EFF6FF", "#2563EB"),
        ((1070, 775, 1370, 910), "Logistik", "Ranking rute, ETA/jarak, estimasi volume, dan kebutuhan data mitra", "#FEF2F2", "#DC2626"),
        ((1405, 775, 1705, 910), "Evidence Room", "Bukti demand, status data, sumber, dan keterbatasan yang transparan", "#F8FAFC", "#0F766E"),
    ]
    for box in output_boxes:
        draw_box(draw, *box)

    spine(925, 960)

    section_label("Pengguna dan Keputusan", 960)
    draw_box(
        draw,
        (210, 990, 1590, 1095),
        "Pengguna Utama",
        "TPID, pemda, dinas pangan, Bapanas, Bulog, BI regional, gudang, carrier, dan mitra distribusi pangan.",
        "#FFFBEB",
        "#D97706",
    )

    draw.text((210, 1120), "Legenda status data:", fill="#334155", font=label_font)
    legend = [
        ("real-time", "#16A34A"),
        ("official-release", "#2563EB"),
        ("forecast", "#7C3AED"),
        ("unavailable / expected partner data", "#DC2626"),
    ]
    x = 425
    for text, color in legend:
        draw.rounded_rectangle((x, 1120, x + 22, 1142), radius=5, fill=color)
        draw.text((x + 32, 1117), text, fill="#475569", font=note_font)
        x += 285

    img.save(path)


def add_footer(doc):
    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("Kepang AI - Submission Attachment - Prepared 4 June 2026")
        run.font.name = "Calibri"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string(MUTED)


def build_doc():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    arch_path = ASSET_DIR / "system_architecture.png"
    create_architecture_image(arch_path)

    doc = Document()
    configure_styles(doc)

    # Cover
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Kepang AI")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(30)
    run.font.color.rgb = RGBColor.from_string(GREEN)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Decision Intelligence untuk Resiliensi Ketahanan Pangan Daerah")
    run.font.name = "Calibri"
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor.from_string(DARK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Combined Submission Attachment - Tahap 2")
    run.font.name = "Calibri"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string(MUTED)

    add_callout(
        doc,
        "Purpose",
        "One PDF package for judges: problem evidence, solution fit, architecture, data policy, impact KPI, business model, MVP readiness, and validation gaps.",
        fill=LIGHT_GREEN,
        color=GREEN,
    )

    add_table(
        doc,
        ["Field", "Content"],
        [
            ["Team ID", "P0684"],
            ["Team Name", "J4"],
            ["Proposal Title", "Kepang AI: Decision Intelligence untuk Resiliensi Ketahanan Pangan Daerah"],
            ["Prototype Status", "Local web prototype with backend API, Supabase database, source-aware forecasting, logistics recommendation, and Evidence Room"],
            ["File Naming Rule", "P0684 - Kepang AI Decision Intelligence untuk Resiliensi Ketahanan Pangan Daerah.pdf"],
        ],
        [1.45, 4.9],
        font_size=9,
    )

    doc.add_page_break()

    add_heading(doc, "1. Executive Evidence Summary")
    add_body(
        doc,
        "Kepang AI is a decision intelligence prototype for food resilience. The core problem is policy lag: price, weather, stock, logistics, and macro pressure are often analyzed separately, so intervention can arrive after food prices have already moved.",
    )
    add_body(
        doc,
        "The solution connects official data, real-time-capable feeds, model-derived forecasts, and explicit data gaps into one decision cockpit. The app turns signals into Resilience Score, scenario shock, route recommendation, action cards, KPI targets, and data lineage.",
    )
    add_table(
        doc,
        ["Evidence Signal", "Why It Matters", "Kepang AI Response"],
        [
            ["Volatile food 6.24% yoy", "Food commodities remain a current inflation pressure", "Food Inflation Monitor and commodity risk"],
            ["USD/IDR Rp17,700", "Weaker rupiah raises imported inflation and logistics exposure", "Rupiah scenario planning and import exposure"],
            ["Weather risk", "Flood, drought, and climate shocks affect production", "BMKG-based Weather Risk Engine"],
            ["Regional imbalance", "Surplus and deficit regions need coordinated distribution", "Supply-demand map and route recommendation"],
            ["Data trust gap", "Forecast, seed, and official data must not be mixed", "Data lineage labels and expected data fields"],
        ],
        [1.6, 2.25, 2.5],
    )

    add_heading(doc, "2. Evidence of Demand")
    add_body(
        doc,
        "Demand is demonstrated by public problem evidence and institutional need. Stakeholders do not only need another dashboard; they need a shared decision layer that explains what happened, why it matters, what action should be prioritized, and which data can be trusted.",
    )
    add_table(
        doc,
        ["Real Problem Evidence", "Source", "Demand Signal"],
        [
            ["Inflasi Mei 2026 mencapai 3.08% yoy dan 0.28% mtm", "BPS, 2 June 2026", "Price stability is an active public issue"],
            ["Volatile food mencapai 6.24% yoy", "Antara citing BPS", "Strategic food commodities still drive inflation pressure"],
            ["Rupiah Rp17,700/USD and BI-Rate 5.25%", "Bank Indonesia RDG May 2026", "Macro shock requires resilient scenario planning"],
            ["Imports Jan-Feb 2026 reached US$42.09B, +14.44% yoy", "BPS export-import release", "Imported inputs, energy, and logistics can become more expensive"],
            ["Bapanas mission includes reserve, logistics, distribution, stabilization, and vulnerable regions", "Bapanas Renstra 2025-2029", "The solution aligns with institutional mandate"],
            ["BMKG provides 3-day village-level weather forecast, updated twice daily", "BMKG Open Data", "Weather data can feed early warning for food risk"],
        ],
        [2.45, 1.45, 2.45],
        font_size=7.8,
    )

    add_heading(doc, "3. Problem-Solution Mapping")
    add_table(
        doc,
        ["Problem", "Mechanism", "Outcome"],
        [
            ["Fragmented food data", "Integrate price, supply-demand, weather, macro, and logistics in one cockpit", "Decision makers see root causes, not only price symptoms"],
            ["Policy lag", "Early warning, Resilience Score, action cards, redistribution recommendation", "Stock intervention, market operation, and harvest mitigation can start earlier"],
            ["Weaker rupiah and imported inflation", "Currency scenario planning, import exposure, logistics-cost simulation", "Food decisions become more resilient against external shocks"],
            ["Unclear data quality", "Data lineage: real-time, official-release, forecast, unavailable", "Every recommendation is more auditable and honest"],
            ["Suboptimal logistics route", "Route ranking, volume estimate, optional Google ETA/distance, expected partner data", "Distribution from surplus to deficit regions becomes more targeted"],
        ],
        [1.85, 2.35, 2.15],
        font_size=7.8,
    )

    doc.add_page_break()

    add_heading(doc, "4. System Architecture")
    add_body(
        doc,
        "The architecture separates user interface, API orchestration, data storage, external data connectors, model outputs, and data lineage. This supports a transparent MVP today and a production path with institutional data partners later.",
    )
    doc.add_picture(str(arch_path), width=Inches(6.35))

    add_heading(doc, "5. Data Source and Feasibility Matrix")
    add_table(
        doc,
        ["Dataset", "Current Status", "Source / Method", "Production Requirement"],
        [
            ["Inflation and macro", "official-release", "BPS and BI releases", "Periodic official update ingestion"],
            ["Weather forecast", "real-time capable", "BMKG Open Data", "adm4 mapping for all target regions"],
            ["Food prices", "real-time capable", "BI Harga Pangan scraper/API; Bapanas when stable", "Stable endpoint access and source logging"],
            ["Supply-demand", "forecast / seed MVP", "Prototype seed; expected Bapanas/Bulog/Dinas data", "production_ton, stock_ton, demand_ton, warehouse_id"],
            ["Logistics ETA/distance", "real-time capable when key exists", "Google Routes API for eligible road routes", "Google key and road-routable origin/destination"],
            ["Logistics cost/capacity", "forecast / unavailable", "Prototype route table", "carrier, capacity_ton, cost_per_ton, lead_time_days"],
            ["Resilience score", "forecast", "Kepang AI model", "Validation against historical interventions"],
        ],
        [1.35, 1.3, 2.0, 1.7],
        font_size=7.4,
    )

    add_callout(
        doc,
        "Data honesty rule",
        "Kepang AI must not present seed, proxy, or model output as audited real-time data. Every displayed dataset is labelled as real-time, official-release, forecast, or unavailable.",
        fill=LIGHT_AMBER,
        color=AMBER,
    )

    doc.add_page_break()

    add_heading(doc, "6. Impact Measurement")
    add_table(
        doc,
        ["KPI", "MVP Target", "Measurement Method"],
        [
            ["Time to identify priority region", "Under 3 minutes", "User task test with TPID/dinas pangan persona"],
            ["Time-to-insight reduction", "50% vs manual baseline", "Compare dashboard task vs spreadsheet/manual workflow"],
            ["Simulated supply gap reduction", "At least 30%", "Before-after redistribution scenario"],
            ["Logistics cost potential", "5-10%", "Cost per ton before-after route optimization"],
            ["Data lineage completeness", "100%", "Audit UI and API response"],
            ["Recommendation validation", "5-10 interviews", "Structured interview scorecard"],
        ],
        [2.05, 1.5, 2.8],
        font_size=8.1,
    )

    add_heading(doc, "7. Business Model and Sustainability")
    add_table(
        doc,
        ["Revenue Stream", "Buyer", "Value Sold"],
        [
            ["Annual institutional license", "Pemda, TPID, Bapanas, Bulog, BI regional", "Dashboard, user seats, modules, regional coverage"],
            ["Implementation fee", "Pilot institution", "Setup, data integration, training, custom workflow"],
            ["Managed analytics", "Government or enterprise partner", "Monthly risk report and scenario briefing"],
            ["API subscription", "Institutions, logistics, finance, insurance", "Risk score, anomaly signal, route intelligence"],
            ["Logistics add-on", "Carrier, warehouse, commodity aggregator", "Route optimization, backhaul, capacity matching"],
            ["Grant / innovation funding", "Public program, CSR, donor", "Early pilot, validation, and social impact deployment"],
        ],
        [1.75, 2.1, 2.5],
        font_size=7.9,
    )
    add_body(
        doc,
        "The product monetizes the intelligence layer, not raw public data. Profitability can come from recurring institutional subscriptions, implementation services, managed analytics, and API/logistics add-ons once pilots prove measurable value.",
    )

    doc.add_page_break()

    add_heading(doc, "8. Guidebook Alignment")
    add_table(
        doc,
        ["Guidebook Criterion", "Evidence in Kepang AI", "Status"],
        [
            ["Alignment with Problem Statement", "Digitalisasi Ketahanan Pangan, policy lag, supply-demand, price, weather, logistics, and rupiah", "Ready"],
            ["Effectiveness & Impact", "Resilience Score, action cards, KPI targets, supply gap, route efficiency", "Ready for MVP"],
            ["Uniqueness / Creativity", "Decision cockpit, rupiah shock scenario, data lineage, source-aware logistics", "Ready"],
            ["Technical Quality", "React/Vite, Express, PostgreSQL/Supabase, BI scraper, BMKG service, Google optional", "Prototype ready"],
            ["Business Model Feasibility", "B2G SaaS, implementation fee, managed analytics, API subscription, logistics add-on", "Ready"],
        ],
        [1.8, 3.55, 1.0],
        font_size=7.7,
    )

    add_heading(doc, "9. MVP Readiness and Test Evidence")
    add_table(
        doc,
        ["Check", "Result", "Observation"],
        [
            ["Frontend build", "Pass", "Vite production build completed; chunk-size warning non-blocking"],
            ["Backend syntax checks", "Pass", "forecast route, logistics route, googleRoutesService valid"],
            ["GET /api/health", "Pass", "status=ok, database=ok, redis=disabled locally"],
            ["GET /api/forecast/resilience", "Pass", "Macro context, resilience score, data policy, data provenance returned"],
            ["GET /api/logistics/recommendations", "Pass", "Ranked routes, source policy, expected production data returned"],
            ["Browser checks", "Pass", "Evidence Room and Dashboard render with no mojibake"],
        ],
        [1.65, 0.9, 3.8],
        font_size=7.8,
    )

    add_heading(doc, "10. Open Gaps and Next Validation")
    add_bullets(
        doc,
        [
            "Fill final Team ID and confirm official team composition before submission.",
            "Run 5-10 structured interviews with TPID, dinas pangan, Bulog/logistics, BI regional, commodity traders, and warehouse/carrier operators.",
            "Integrate production stock, warehouse, carrier capacity, port congestion, and actual cost-per-ton feeds when partner access is available.",
            "Configure Google Maps keys only for road ETA/distance; do not claim sea/cargo cost or capacity as Google real-time data.",
            "Prepare a 2-3 minute demo video as backup if public hosting is deferred.",
        ]
    )

    add_heading(doc, "11. Source Links")
    source_rows = [
        ["BPS Inflasi Mei 2026", "https://www.bps.go.id/id/pressrelease/2026/06/02/2579/inflasi-year-on-year--y-on-y--pada-mei-2026-sebesar-3-08-persen-.html"],
        ["Antara/BPS volatile food", "https://www.antaranews.com/berita/5590715/bps-kenaikan-bbm-nonsubsidi-dorong-inflasi-transportasi-061-persen"],
        ["BI RDG Mei 2026", "https://www.bi.go.id/id/publikasi/ruang-media/news-release/Pages/sp_2810726.aspx"],
        ["BPS Ekspor-Impor Februari 2026", "https://www.bps.go.id/assets/pressrelease/2026/04/01/2557/ekspor-dan-impor-indonesia-februari-2026-masing-masing-tercatat-usd-22-17-miliar-dan-usd-20-89-miliar-.html"],
        ["BPS NTP Mei 2026", "https://www.bps.go.id/id/pressrelease/2026/06/02/2580/nilai-tukar-petani--ntp--mei-2026-sebesar-127-73-atau-naik-1-99-persen.html"],
        ["BMKG Open Data", "https://data.bmkg.go.id/prakiraan-cuaca"],
        ["Bapanas Renstra 2025-2029", "https://peraturan.go.id/filespengundangan/peraturan-bapanas-no-9-tahun-2025.pdf"],
    ]
    add_table(doc, ["Source", "URL"], source_rows, [1.85, 4.5], font_size=6.8)

    add_footer(doc)
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_doc()
